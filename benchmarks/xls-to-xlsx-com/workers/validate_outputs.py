"""Validate benchmark XLSX outputs outside the timed conversion window."""

from __future__ import annotations

import argparse
import json
import math
import sys
from pathlib import Path

import openpyxl


def fail(message: str) -> None:
    raise AssertionError(message)


def close_enough(actual: float, expected: float, tolerance: float = 0.75) -> bool:
    return math.isfinite(actual) and abs(actual - expected) <= tolerance


def validate_one(path: Path, expected: dict) -> list[str]:
    errors: list[str] = []
    if not path.is_file():
        return [f"missing output: {path.name}"]
    workbook = None
    try:
        workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
        expected_sheets = expected["sheets"]
        if workbook.sheetnames != [sheet["name"] for sheet in expected_sheets]:
            errors.append(f"{path.name}: sheet names/order mismatch")
        if len(workbook.worksheets) != len(expected_sheets):
            errors.append(f"{path.name}: sheet count mismatch")

        first_sheet = workbook.worksheets[0]
        for cell, value in expected["key_values"].items():
            if first_sheet[cell].value != value:
                errors.append(f"{path.name}: key value {cell} mismatch")

        expected_visibility = {"visible": "visible", "hidden": "hidden", "veryHidden": "veryHidden"}
        for expected_sheet in expected_sheets:
            if expected_sheet["name"] not in workbook.sheetnames:
                continue
            sheet = workbook[expected_sheet["name"]]
            if sheet.sheet_state != expected_visibility[expected_sheet["hidden"]]:
                errors.append(f"{path.name}/{sheet.title}: visibility mismatch")
            if expected_sheet["merge"] not in {str(value) for value in sheet.merged_cells.ranges}:
                errors.append(f"{path.name}/{sheet.title}: merged range mismatch")
            formula_cell = expected_sheet["formula_cell"]
            if sheet[formula_cell].value != expected_sheet["formula"]:
                errors.append(f"{path.name}/{sheet.title}: formula text mismatch at {formula_cell}")
            if not sheet["A1"].font.bold or sheet["A1"].fill.fill_type is None:
                errors.append(f"{path.name}/{sheet.title}: title key style missing")
            if not sheet["A2"].font.bold or sheet["A2"].fill.fill_type is None:
                errors.append(f"{path.name}/{sheet.title}: header key style missing")
            for cell, style in expected_sheet["style_checks"].items():
                if "number_format" in style and sheet[cell].number_format != style["number_format"]:
                    errors.append(f"{path.name}/{sheet.title}: number format mismatch at {cell}")
                if style.get("custom_number_format") and sheet[cell].number_format == "General":
                    errors.append(f"{path.name}/{sheet.title}: custom number format missing at {cell}")
            for row, expected_height in expected_sheet["row_heights"].items():
                actual_height = sheet.row_dimensions[int(row)].height
                if actual_height is None or not close_enough(float(actual_height), float(expected_height)):
                    errors.append(f"{path.name}/{sheet.title}: row {row} height mismatch")
            for column, expected_width in expected_sheet["column_widths"].items():
                actual_width = sheet.column_dimensions[column].width
                if actual_width is None or not close_enough(float(actual_width), float(expected_width)):
                    errors.append(f"{path.name}/{sheet.title}: column {column} width mismatch")

        defined_names = set(workbook.defined_names)
        if expected["defined_name"] not in defined_names:
            errors.append(f"{path.name}: named range {expected['defined_name']} missing")
    except Exception as error:
        errors.append(f"{path.name}: reopen/validation failed: {type(error).__name__}: {error}")
    finally:
        if workbook is not None:
            workbook.close()
    return errors


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--manifest", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--files", required=True, type=Path, help="JSON list of source filenames")
    args = parser.parse_args()
    manifest = json.loads(args.manifest.read_text(encoding="utf-8"))
    selected = json.loads(args.files.read_text(encoding="utf-8"))
    expected_by_name = {item["name"]: item for item in manifest["items"]}
    expected_names = [Path(name).with_suffix(".xlsx").name for name in selected]
    actual_names = sorted(path.name for path in args.output_dir.glob("*.xlsx"))
    errors = []
    if sorted(expected_names) != actual_names:
        errors.append(
            f"output file count/name mismatch: expected {len(expected_names)}, found {len(actual_names)}"
        )
    for source_name in selected:
        if source_name not in expected_by_name:
            errors.append(f"source {source_name} is absent from the fixed manifest")
            continue
        output_name = Path(source_name).with_suffix(".xlsx").name
        errors.extend(validate_one(args.output_dir / output_name, expected_by_name[source_name]))

    result = {
        "validation": "passed" if not errors else "failed",
        "expected_output_count": len(expected_names),
        "actual_output_count": len(actual_names),
        "errors": errors,
    }
    print(json.dumps(result, ensure_ascii=False, separators=(",", ":")))
    return 0 if not errors else 1


if __name__ == "__main__":
    raise SystemExit(main())
